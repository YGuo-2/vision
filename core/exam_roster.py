# -*- coding: utf-8 -*-
"""考试名单导入、成绩台账与 Excel 导出。

分数口径：内部 front/side/combined_score 为 0..1；导出时正侧 ×100（一位小数），
综合用 combined_percent（0–100 整数）。
"""

from __future__ import annotations

import threading
import time
from dataclasses import dataclass, field
from datetime import datetime
from pathlib import Path
from typing import Any, Callable, Iterable, Literal

IMPORT_HEADERS = ("序号", "学号", "姓名", "班级", "备注")
EXPORT_HEADERS = (
    "序号",
    "学号",
    "姓名",
    "班级",
    "综合分",
    "正面分",
    "侧面分",
    "状态",
    "错误信息",
    "片段ID",
    "录制目录",
    "叫号时间",
    "开录时间",
    "结束时间",
    "备注",
    "尝试次数",
    "是否计分行",
)

RowStatus = Literal[
    "pending",
    "recording",
    "processing",
    "completed",
    "failed",
    "skipped",
    "superseded",
]


class RosterError(ValueError):
    def __init__(self, code: str, message: str) -> None:
        super().__init__(message)
        self.code = code
        self.message = message


@dataclass(frozen=True)
class ExamCandidate:
    order: int
    student_id: str
    name: str
    class_name: str = ""
    note: str = ""


@dataclass
class ExamResultRow:
    candidate: ExamCandidate
    status: RowStatus = "pending"
    segment_id: str | None = None
    segment_dir: str | None = None
    front_score: float | None = None  # 0..1
    side_score: float | None = None  # 0..1
    combined_score: float | None = None  # 0..1
    combined_percent: int | None = None  # 0..100
    error_code: str | None = None
    error_message: str | None = None
    warnings: list[str] = field(default_factory=list)
    # 逐动作子分：动作名 → 0..1 相似度分（宽表导出为「<动作>分」列，×100 一位小数）
    action_scores: dict[str, float] = field(default_factory=dict)
    called_at: str | None = None
    record_started_at: str | None = None
    record_ended_at: str | None = None
    attempt_index: int = 1
    is_scoring_row: bool = True
    row_id: str = ""

    def __post_init__(self) -> None:
        if not self.row_id:
            self.row_id = (
                f"{self.candidate.student_id}#{self.attempt_index}"
                f"@{self.candidate.order}"
            )


def local_timestamp() -> str:
    return datetime.now().astimezone().isoformat(timespec="seconds")


def _cell_text(value: Any) -> str:
    if value is None:
        return ""
    if isinstance(value, bool):
        return "true" if value else "false"
    if isinstance(value, int):
        return str(value)
    if isinstance(value, float):
        if value.is_integer() and abs(value) < 1e15:
            return str(int(value))
        # 科学计数法风险：不自动“还原”学号，交给上层校验
        text = format(value, ".15g")
        return text
    return str(value).strip()


def _excel_format_is_plain_number(number_format: str | None) -> bool:
    """General / @ / 0 等视为「无自定义显示」，可安全回退为裸数字。"""
    import re

    fmt = (number_format or "").strip()
    if not fmt or fmt in {"General", "@", "0"}:
        return True
    first = fmt.split(";")[0].strip()
    first = re.sub(r"\[[^\]]*\]", "", first).strip()
    return first in {"", "General", "@", "0"}


def _tokenize_excel_number_format(number_format: str | None) -> list[tuple[str, str]] | None:
    """解析 Excel 数字格式第一段为 token 列表：('d','0'|'#') 或 ('lit', ch)。

    支持反斜杠转义（``000\\-000``）与双引号字面量（``\"ID-\"000000``）。
    不支持小数/千分位/科学计数/条件格式表达式时返回 None。
    """
    import re

    fmt = (number_format or "").strip()
    if not fmt:
        return None
    first = fmt.split(";")[0]
    first = re.sub(r"\[[^\]]*\]", "", first)
    if not first.strip() or first.strip() in {"General", "@"}:
        return None

    tokens: list[tuple[str, str]] = []
    i = 0
    n = len(first)
    while i < n:
        ch = first[i]
        if ch == "\\":
            if i + 1 >= n:
                return None
            tokens.append(("lit", first[i + 1]))
            i += 2
            continue
        if ch == '"':
            i += 1
            buf: list[str] = []
            while i < n:
                if first[i] == '"':
                    if i + 1 < n and first[i + 1] == '"':
                        buf.append('"')
                        i += 2
                        continue
                    i += 1
                    break
                buf.append(first[i])
                i += 1
            else:
                return None  # 未闭合引号
            for c in buf:
                tokens.append(("lit", c))
            continue
        if ch in "0#":
            tokens.append(("d", ch))
            i += 1
            continue
        if ch in "- /()_:.":
            # 字面分隔；小数点单独拒绝（与科学/会计格式混淆）
            if ch == ".":
                return None
            tokens.append(("lit", ch))
            i += 1
            continue
        if ch in {",", "E", "e", "%", "*", "?", "_"}:
            return None
        # 其他字符（含中文）视为字面量，便于少量自定义
        if ch.isalpha() or ord(ch) > 127:
            tokens.append(("lit", ch))
            i += 1
            continue
        return None
    if not any(t[0] == "d" for t in tokens):
        return None
    return tokens


def _display_text_from_number_format(value: int | float, number_format: str | None) -> str | None:
    """尽量按 Excel 显示文本还原学号。

    支持：
    - 纯补零：``000000`` + 123 → ``000123``
    - ``0`` 必选位 / ``#`` 可选位：``###000`` + 123 → ``123``（不按 # 强制补零）
    - 分隔：``000-000`` / ``000\\-000`` + 123 → ``000-123``
    - 字面量前缀：``\"ID-\"000000`` + 123 → ``ID-000123``
    - 超出占位宽度时保留字面量：``\"ID-\"000`` + 12345 → ``ID-12345``

    无法安全解析时返回 None（调用方对自定义格式应整表拒绝，不得静默裸数字）。
    """
    tokens = _tokenize_excel_number_format(number_format)
    if tokens is None:
        if _excel_format_is_plain_number(number_format):
            return str(int(value)) if float(value).is_integer() else None
        return None
    ph_types = [t[1] for t in tokens if t[0] == "d"]
    if not ph_types:
        return None
    n = int(value)
    if n < 0:
        return None
    # 仅按必选 0 位补零；# 不强制加宽
    min_len = sum(1 for p in ph_types if p == "0")
    # Excel：纯 #（无必选 0 位）对数值 0 显示为空，不得导入成 "0"
    if n == 0 and min_len == 0:
        digit_str = ""
    else:
        digit_str = str(n)
        if len(digit_str) < min_len:
            digit_str = digit_str.zfill(min_len)

    # 从右向左填入占位；多余数字保留为 excess（仍挂在字面量之后）
    slots: list[str | None] = [None] * len(ph_types)
    di = len(digit_str) - 1
    for si in range(len(ph_types) - 1, -1, -1):
        if di >= 0:
            slots[si] = digit_str[di]
            di -= 1
        elif ph_types[si] == "0":
            slots[si] = "0"
        else:
            slots[si] = None  # 未使用的 #
    excess = digit_str[: di + 1] if di >= 0 else ""

    out: list[str] = []
    si = 0
    excess_pending = excess
    for kind, payload in tokens:
        if kind == "d":
            ch = slots[si]
            si += 1
            if ch is None:
                continue
            if excess_pending:
                out.append(excess_pending)
                excess_pending = ""
            out.append(ch)
        else:
            out.append(payload)
    if excess_pending:
        # 无数字位可挂时（极端），前缀 excess
        out.insert(0, excess_pending)
    result = "".join(out)
    # 零值 + 仅 # 位：显示无数字 → 空串（调用方按空学号拒绝）
    if n == 0 and min_len == 0 and not any(c.isdigit() for c in result):
        return ""
    return result


def _student_id_from_cell(value: Any, *, row_no: int, number_format: str | None = None) -> str:
    """强制按文本语义读取学号；数值单元格优先用 number_format 显示文本。

    自定义格式无法还原时 **拒绝**（不静默变成裸数字），迫使将学号列设为文本。
    """
    if value is None:
        raise RosterError("empty_student_id", f"第 {row_no} 行学号为空")
    if isinstance(value, str):
        text = value.strip()
        if not text:
            raise RosterError("empty_student_id", f"第 {row_no} 行学号为空")
        return text
    if isinstance(value, bool):
        raise RosterError(
            "student_id_not_text",
            f"第 {row_no} 行学号类型无效：{value!r}",
        )
    if isinstance(value, int):
        display = _display_text_from_number_format(value, number_format)
        if display is not None:
            if not str(display).strip():
                raise RosterError(
                    "empty_student_id",
                    f"第 {row_no} 行学号为空（显示格式下无有效文本，如 0+###）",
                )
            return display
        if _excel_format_is_plain_number(number_format):
            return str(value)
        raise RosterError(
            "student_id_not_text",
            f"第 {row_no} 行学号无法按显示格式 {number_format!r} 还原，"
            "请将该列设为文本格式后重试",
        )
    if isinstance(value, float):
        # 允许整数值 float（Excel 有时如此），但拒绝非整数 / 过大科学计数
        if not value.is_integer():
            raise RosterError(
                "student_id_not_text",
                f"第 {row_no} 行学号疑似被 Excel 读成小数/科学计数法，请将该列设为文本后重试",
            )
        if abs(value) >= 1e15:
            raise RosterError(
                "student_id_not_text",
                f"第 {row_no} 行学号过大，请将该列设为文本格式后重试",
            )
        ival = int(value)
        display = _display_text_from_number_format(ival, number_format)
        if display is not None:
            if not str(display).strip():
                raise RosterError(
                    "empty_student_id",
                    f"第 {row_no} 行学号为空（显示格式下无有效文本，如 0+###）",
                )
            return display
        if _excel_format_is_plain_number(number_format):
            return str(ival)
        raise RosterError(
            "student_id_not_text",
            f"第 {row_no} 行学号无法按显示格式 {number_format!r} 还原，"
            "请将该列设为文本格式后重试",
        )
    text = _cell_text(value)
    if not text:
        raise RosterError("empty_student_id", f"第 {row_no} 行学号为空")
    return text


def import_roster_xlsx(path: Path | str) -> list[ExamCandidate]:
    """导入第一个 sheet；失败整表拒绝。"""
    try:
        from openpyxl import load_workbook
    except ImportError as exc:
        raise RosterError("openpyxl_missing", "缺少 openpyxl，请先安装依赖") from exc

    path = Path(path)
    if not path.is_file():
        raise RosterError("file_missing", f"名单文件不存在：{path}")

    try:
        # data_only=False：保留 number_format，便于学号按「显示文本」读取（如 000123）
        wb = load_workbook(path, read_only=False, data_only=False)
    except Exception as exc:
        raise RosterError("file_unreadable", f"无法读取 Excel：{exc}") from exc

    try:
        ws = wb[wb.sheetnames[0]]
        if getattr(ws, "merged_cells", None) and ws.merged_cells.ranges:
            raise RosterError(
                "merged_cells",
                "导入表存在合并单元格，请拆分后再导入",
            )

        rows_iter = ws.iter_rows()
        try:
            header_cells = next(rows_iter)
        except StopIteration as exc:
            raise RosterError("empty_sheet", "Excel 为空") from exc

        headers = [_cell_text(c.value) for c in header_cells]
        # 去掉尾部空表头
        while headers and not headers[-1]:
            headers.pop()
        index = {name: i for i, name in enumerate(headers) if name}
        for required in ("序号", "学号", "姓名"):
            if required not in index:
                raise RosterError("missing_column", f"缺少必填列：{required}")

        candidates: list[ExamCandidate] = []
        seen_orders: set[int] = set()
        for excel_row_no, row_cells in enumerate(rows_iter, start=2):
            cells = list(row_cells)
            while len(cells) < len(headers):
                cells.append(None)  # type: ignore[arg-type]
            values = [c.value if c is not None else None for c in cells[: len(headers)]]
            # 全空行跳过
            if all(v is None or _cell_text(v) == "" for v in values):
                continue

            order_raw = values[index["序号"]]
            try:
                if isinstance(order_raw, float) and order_raw.is_integer():
                    order = int(order_raw)
                else:
                    order = int(str(order_raw).strip())
            except (TypeError, ValueError) as exc:
                raise RosterError(
                    "bad_order", f"第 {excel_row_no} 行序号无效：{order_raw!r}"
                ) from exc
            if order <= 0:
                raise RosterError("bad_order", f"第 {excel_row_no} 行序号须为正整数")
            if order in seen_orders:
                raise RosterError("duplicate_order", f"序号重复：{order}")
            seen_orders.add(order)

            sid_cell = cells[index["学号"]]
            sid_fmt = getattr(sid_cell, "number_format", None) if sid_cell is not None else None
            student_id = _student_id_from_cell(
                values[index["学号"]],
                row_no=excel_row_no,
                number_format=sid_fmt,
            )
            name = _cell_text(values[index["姓名"]])
            if not name:
                raise RosterError("empty_name", f"第 {excel_row_no} 行姓名为空")

            class_name = (
                _cell_text(values[index["班级"]]) if "班级" in index else ""
            )
            note = _cell_text(values[index["备注"]]) if "备注" in index else ""
            candidates.append(
                ExamCandidate(
                    order=order,
                    student_id=student_id,
                    name=name,
                    class_name=class_name,
                    note=note,
                )
            )

        if not candidates:
            raise RosterError("no_rows", "没有有效考生行")
        candidates.sort(key=lambda c: c.order)
        return candidates
    finally:
        wb.close()


def write_import_template(path: Path | str) -> Path:
    """写出导入模板（表头 + 示例行，学号列文本格式）。"""
    from openpyxl import Workbook
    from openpyxl.styles import numbers

    path = Path(path)
    path.parent.mkdir(parents=True, exist_ok=True)
    wb = Workbook()
    ws = wb.active
    ws.title = "名单"
    ws.append(list(IMPORT_HEADERS))
    ws.append([1, "2026001", "张三", "一年级1班", ""])
    ws.append([2, "2026002", "李四", "一年级1班", ""])
    for row in range(2, 4):
        ws.cell(row=row, column=2).number_format = numbers.FORMAT_TEXT
    wb.save(path)
    wb.close()
    return path


def _format_front_side_export(score: float | None) -> str | float:
    if score is None:
        return ""
    return round(float(score) * 100.0, 1)


def collect_action_names(rows: Iterable[ExamResultRow]) -> list[str]:
    """按各行首次出现顺序收集所有动作名（去重保序）→ 宽表动作列顺序。"""
    names: list[str] = []
    seen: set[str] = set()
    for row in rows:
        for name in row.action_scores:
            if name not in seen:
                seen.add(name)
                names.append(name)
    return names


def row_to_export_values(
    row: ExamResultRow, action_names: list[str] | None = None
) -> list[Any]:
    err = ""
    if row.error_code or row.error_message:
        err = f"{row.error_code or ''}: {row.error_message or ''}".strip(": ")
    values: list[Any] = [
        row.candidate.order,
        row.candidate.student_id,
        row.candidate.name,
        row.candidate.class_name,
        row.combined_percent if row.combined_percent is not None else "",
        _format_front_side_export(row.front_score),
        _format_front_side_export(row.side_score),
        row.status,
        err,
        row.segment_id or "",
        row.segment_dir or "",
        row.called_at or "",
        row.record_started_at or "",
        row.record_ended_at or "",
        row.candidate.note,
        row.attempt_index,
        "是" if row.is_scoring_row else "否",
    ]
    # 动作分列追加表尾：×100 一位小数，缺失动作留空
    for name in action_names or []:
        values.append(_format_front_side_export(row.action_scores.get(name)))
    return values


def write_scorebook_xlsx(path: Path | str, rows: Iterable[ExamResultRow]) -> None:
    """同步写出成绩表（原子 replace）。调用方应处理 PermissionError。"""
    from openpyxl import Workbook

    path = Path(path)
    path.parent.mkdir(parents=True, exist_ok=True)
    rows = list(rows)
    action_names = collect_action_names(rows)
    wb = Workbook()
    ws = wb.active
    ws.title = "成绩"
    ws.append(list(EXPORT_HEADERS) + [f"{n}分" for n in action_names])
    for row in rows:
        ws.append(row_to_export_values(row, action_names))

    # 临时文件用 ASCII 名，避免部分环境下中文文件名 + 前导点的怪异行为
    tmp = path.parent / f"_scorebook_{time.time_ns()}.tmp.xlsx"
    try:
        wb.save(tmp)
        wb.close()
        _atomic_replace(tmp, path)
    finally:
        if tmp.exists():
            try:
                tmp.unlink()
            except OSError:
                pass


def _atomic_replace(src: Path, dst: Path) -> None:
    """同卷 replace；失败时 copy+replace 回退。"""
    import shutil

    try:
        src.replace(dst)
        return
    except OSError:
        pass
    # 跨盘或部分 Windows 场景
    bak = dst.with_suffix(dst.suffix + ".bak")
    try:
        if dst.exists():
            try:
                dst.replace(bak)
            except OSError:
                bak = None  # type: ignore[assignment]
        shutil.copy2(src, dst)
        src.unlink(missing_ok=True)
        if bak is not None and Path(bak).exists():
            Path(bak).unlink(missing_ok=True)
    except OSError:
        if bak is not None and Path(bak).exists() and not dst.exists():
            try:
                Path(bak).replace(dst)
            except OSError:
                pass
        raise


class ExamScorebook:
    """考试成绩台账：行列表 + 可选异步落盘。"""

    def __init__(
        self,
        *,
        path: Path | None = None,
        write_fn: Callable[[Path, list[ExamResultRow]], None] | None = None,
        throttle_s: float = 2.0,
    ) -> None:
        self.path = Path(path) if path else None
        self._write_fn = write_fn or write_scorebook_xlsx
        self._throttle_s = float(throttle_s)
        self._lock = threading.RLock()
        # 串行化落盘，配合 generation 防止旧快照覆盖新快照
        self._write_lock = threading.Lock()
        self._rows: list[ExamResultRow] = []
        self._by_segment: dict[str, ExamResultRow] = {}
        self._dirty = False
        self._generation = 0
        self._written_generation = 0
        self._last_write = 0.0
        self._write_error: str | None = None
        self._closed = False
        self._worker: threading.Thread | None = None
        self._wake = threading.Event()
        if self.path is not None:
            self._worker = threading.Thread(
                target=self._writer_loop, name="exam-scorebook-writer", daemon=True
            )
            self._worker.start()

    @property
    def rows(self) -> list[ExamResultRow]:
        with self._lock:
            return list(self._rows)

    @property
    def write_error(self) -> str | None:
        with self._lock:
            return self._write_error

    def load_candidates(self, candidates: list[ExamCandidate]) -> None:
        with self._lock:
            self._rows = [
                ExamResultRow(candidate=c, attempt_index=1, is_scoring_row=True)
                for c in candidates
            ]
            self._by_segment.clear()
            self._mark_dirty_unlocked()

    def append_retest(self, candidate: ExamCandidate) -> ExamResultRow:
        """队尾追加同学号重考行；计分行策略见 ``_apply_scoring_policy_unlocked``。"""
        with self._lock:
            attempts = [
                r.attempt_index
                for r in self._rows
                if r.candidate.student_id == candidate.student_id
            ]
            next_attempt = (max(attempts) if attempts else 0) + 1
            # 已有成功分时，新补考暂不计分，直至其 completed
            has_completed = any(
                r.candidate.student_id == candidate.student_id
                and r.status in {"completed", "superseded"}
                and r.combined_percent is not None
                for r in self._rows
            )
            row = ExamResultRow(
                candidate=candidate,
                attempt_index=next_attempt,
                is_scoring_row=not has_completed,
            )
            self._rows.append(row)
            self._apply_scoring_policy_unlocked(candidate.student_id)
            self._mark_dirty_unlocked()
            return row

    def get_by_segment(self, segment_id: str) -> ExamResultRow | None:
        with self._lock:
            return self._by_segment.get(segment_id)

    def bind_segment(
        self, row: ExamResultRow, *, segment_id: str, segment_dir: str | None
    ) -> None:
        with self._lock:
            if row.segment_id and row.segment_id in self._by_segment:
                self._by_segment.pop(row.segment_id, None)
            row.segment_id = segment_id
            row.segment_dir = segment_dir
            self._by_segment[segment_id] = row
            self._mark_dirty_unlocked()

    def update_row(self, row: ExamResultRow, **fields: Any) -> None:
        with self._lock:
            # 终态保护：
            # - completed/skipped/superseded 不得被 recording/processing 等弱状态盖回
            # - failed 允许回到 recording（BEGIN 失败后同人重试，见 begin_failed 契约）
            _hard_terminal = frozenset(
                {"completed", "skipped", "superseded", "cancelled"}
            )
            new_status = fields.get("status")
            if (
                new_status is not None
                and row.status in _hard_terminal
                and new_status not in _hard_terminal
                and new_status != "failed"
                and new_status != row.status
            ):
                fields = {k: v for k, v in fields.items() if k != "status"}
            for key, value in fields.items():
                if not hasattr(row, key):
                    raise AttributeError(key)
                setattr(row, key, value)
            if row.status in {
                "completed",
                "failed",
                "skipped",
                "superseded",
                "processing",
                "pending",
            }:
                self._apply_scoring_policy_unlocked(row.candidate.student_id)
            self._mark_dirty_unlocked(force=row.status in {
                "completed", "failed", "skipped", "superseded", "processing"
            })

    def _apply_scoring_policy_unlocked(self, student_id: str) -> None:
        """同学号计分行策略：

        - 有成功分：最后一次 completed（含可恢复的 superseded）为唯一计分行；
          更早成功 → superseded 且 is_scoring_row=False。
        - 无成功分：仅 attempt_index 最大的一行 is_scoring_row=True
          （避免「首次成功 + 补考失败」双行同时计分，也避免双失败双计分）。
        """
        rows = [r for r in self._rows if r.candidate.student_id == student_id]
        if not rows:
            return
        completed = [
            r
            for r in rows
            if r.combined_percent is not None
            and r.status in {"completed", "superseded"}
        ]
        if completed:
            completed.sort(key=lambda r: r.attempt_index)
            winner = completed[-1]
            for r in rows:
                if r is winner:
                    r.status = "completed"
                    r.is_scoring_row = True
                else:
                    if r.combined_percent is not None and r.status in {
                        "completed",
                        "superseded",
                    }:
                        r.status = "superseded"
                    r.is_scoring_row = False
            return
        last = max(rows, key=lambda r: r.attempt_index)
        for r in rows:
            r.is_scoring_row = r is last

    def scoring_rows(self) -> list[ExamResultRow]:
        """导出用：每学号优先最后成功，否则最后一次尝试。"""
        with self._lock:
            by_sid: dict[str, list[ExamResultRow]] = {}
            for r in self._rows:
                by_sid.setdefault(r.candidate.student_id, []).append(r)
            out: list[ExamResultRow] = []
            # 保持首次出现顺序
            seen: list[str] = []
            for r in self._rows:
                if r.candidate.student_id not in seen:
                    seen.append(r.candidate.student_id)
            for sid in seen:
                group = by_sid[sid]
                completed = [x for x in group if x.status == "completed"]
                if completed:
                    out.append(max(completed, key=lambda x: x.attempt_index))
                else:
                    out.append(max(group, key=lambda x: x.attempt_index))
            return out

    def flush(self, *, timeout: float = 5.0) -> None:
        """请求立即写盘并等待；同步落盘且不让旧快照覆盖更新代。"""
        with self._lock:
            self._dirty = True
            self._last_write = 0.0
            path = self.path
        self._wake.set()
        if path is None:
            return
        self._persist_latest(force=True)
        _ = timeout  # API 兼容

    def close(self) -> None:
        """关闭台账：有限次终刷后停止后台 writer（文件被锁时不永久重试）。"""
        with self._lock:
            self._closed = True
            self._dirty = True
        self._wake.set()
        # 同步有限次终刷（例如 Excel 占用时最多 2 次）
        if self.path is not None:
            for _ in range(2):
                self._persist_latest(force=True)
                with self._lock:
                    if not self._dirty:
                        break
            # 仍失败则放弃后台重试，保留 write_error 供 UI 提示
            with self._lock:
                self._dirty = False
        worker = self._worker
        if worker is not None and worker.is_alive():
            worker.join(timeout=3.0)

    def _mark_dirty_unlocked(self, *, force: bool = False) -> None:
        self._dirty = True
        self._generation += 1
        if force:
            self._last_write = 0.0
        self._wake.set()

    def _persist_latest(self, *, force: bool = False) -> None:
        """在写锁内取最新快照写盘；仅当 generation 未前进时清 dirty。"""
        with self._write_lock:
            with self._lock:
                if self.path is None:
                    return
                if not self._dirty and self._generation == self._written_generation:
                    return
                if (
                    not force
                    and not self._closed
                    and (time.monotonic() - self._last_write) < self._throttle_s
                    and self._dirty
                ):
                    # 节流：后台循环会再试；force/flush/close 跳过
                    return
                rows_snapshot = list(self._rows)
                path = self.path
                gen = self._generation
                closed = self._closed
            try:
                self._write_fn(path, rows_snapshot)
                with self._lock:
                    if gen >= self._written_generation:
                        self._written_generation = gen
                    # 写盘期间若又有更新，保留 dirty 让后续再刷
                    if gen == self._generation:
                        self._dirty = False
                    self._write_error = None
                    self._last_write = time.monotonic()
            except PermissionError as exc:
                with self._lock:
                    self._write_error = (
                        f"scorebook_locked: 成绩文件被占用，请关闭 Excel 后重试（{exc}）"
                    )
                    # close 之后不再靠 dirty 驱动无限重试
                    self._dirty = not closed
            except OSError as exc:
                with self._lock:
                    self._write_error = f"scorebook_write_failed: {exc}"
                    self._dirty = not closed

    def _writer_loop(self) -> None:
        while True:
            self._wake.wait(timeout=0.5)
            self._wake.clear()
            with self._lock:
                if self._closed and not self._dirty:
                    return
                if self.path is None:
                    if self._closed:
                        return
                    continue
                if not self._dirty and self._generation == self._written_generation:
                    if self._closed:
                        return
                    continue
                now = time.monotonic()
                if (
                    (now - self._last_write) < self._throttle_s
                    and not self._closed
                    and self._dirty
                ):
                    continue
                closed = self._closed
            self._persist_latest(force=closed)
            with self._lock:
                if self._closed:
                    # close 周期只做有限尝试；避免 PermissionError 下永久循环
                    self._dirty = False
                    return
