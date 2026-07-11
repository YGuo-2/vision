# -*- coding: utf-8 -*-
from __future__ import annotations

from pathlib import Path

import pytest

from core.exam_roster import (
    ExamCandidate,
    ExamResultRow,
    ExamScorebook,
    RosterError,
    import_roster_xlsx,
    write_import_template,
    write_scorebook_xlsx,
)


def test_import_template_and_roundtrip(tmp_path: Path) -> None:
    path = tmp_path / "roster.xlsx"
    write_import_template(path)
    cands = import_roster_xlsx(path)
    assert len(cands) == 2
    assert cands[0].student_id == "2026001"
    assert cands[0].name == "张三"


def test_reject_duplicate_order(tmp_path: Path) -> None:
    from openpyxl import Workbook

    path = tmp_path / "bad.xlsx"
    wb = Workbook()
    ws = wb.active
    ws.append(["序号", "学号", "姓名"])
    ws.append([1, "A", "甲"])
    ws.append([1, "B", "乙"])
    wb.save(path)
    wb.close()
    with pytest.raises(RosterError) as ei:
        import_roster_xlsx(path)
    assert ei.value.code == "duplicate_order"


def test_reject_merged_cells(tmp_path: Path) -> None:
    from openpyxl import Workbook

    path = tmp_path / "merged.xlsx"
    wb = Workbook()
    ws = wb.active
    ws.append(["序号", "学号", "姓名"])
    ws.append([1, "A", "甲"])
    ws.merge_cells("A1:B1")
    wb.save(path)
    wb.close()
    with pytest.raises(RosterError) as ei:
        import_roster_xlsx(path)
    assert ei.value.code == "merged_cells"


def test_student_id_float_integer_ok(tmp_path: Path) -> None:
    from openpyxl import Workbook

    path = tmp_path / "sid.xlsx"
    wb = Workbook()
    ws = wb.active
    ws.append(["序号", "学号", "姓名"])
    ws.append([1, 2026001.0, "甲"])
    wb.save(path)
    wb.close()
    cands = import_roster_xlsx(path)
    assert cands[0].student_id == "2026001"


def test_student_id_non_integer_float_rejected(tmp_path: Path) -> None:
    from openpyxl import Workbook

    path = tmp_path / "sid2.xlsx"
    wb = Workbook()
    ws = wb.active
    ws.append(["序号", "学号", "姓名"])
    ws.append([1, 1.5, "甲"])
    wb.save(path)
    wb.close()
    with pytest.raises(RosterError) as ei:
        import_roster_xlsx(path)
    assert ei.value.code == "student_id_not_text"


def test_export_score_scale(tmp_path: Path) -> None:
    path = tmp_path / "scores.xlsx"
    row = ExamResultRow(
        candidate=ExamCandidate(1, "S1", "张三", "一班"),
        status="completed",
        front_score=0.823,
        side_score=0.901,
        combined_score=0.86,
        combined_percent=86,
    )
    write_scorebook_xlsx(path, [row])
    from openpyxl import load_workbook

    wb = load_workbook(path)
    ws = wb.active
    values = [c.value for c in next(ws.iter_rows(min_row=2, max_row=2))]
    assert values[4] == 86
    assert values[5] == 82.3
    assert values[6] == 90.1
    wb.close()


def test_scoring_policy_last_completed_wins() -> None:
    book = ExamScorebook(path=None)
    c = ExamCandidate(1, "S1", "张三")
    book.load_candidates([c])
    r1 = book.rows[0]
    book.update_row(
        r1,
        status="completed",
        front_score=0.5,
        side_score=0.5,
        combined_score=0.5,
        combined_percent=50,
    )
    r2 = book.append_retest(c)
    assert r2.is_scoring_row is False  # 已有成功分时补考暂不计分
    book.update_row(
        r2,
        status="completed",
        front_score=0.9,
        side_score=0.9,
        combined_score=0.9,
        combined_percent=90,
    )
    scoring = book.scoring_rows()
    assert len(scoring) == 1
    assert scoring[0].combined_percent == 90
    assert r1.status == "superseded"
    assert r1.is_scoring_row is False
    assert r2.is_scoring_row is True


def test_update_row_refuses_to_downgrade_hard_terminal_but_allows_failed_retry() -> None:
    book = ExamScorebook(path=None)
    c = ExamCandidate(1, "S1", "甲")
    book.load_candidates([c])
    r = book.rows[0]
    book.update_row(
        r,
        status="completed",
        front_score=0.5,
        side_score=0.5,
        combined_score=0.5,
        combined_percent=50,
    )
    book.update_row(r, status="recording")
    assert r.status == "completed"
    # begin_failed 后允许重试回 recording
    r2_book = ExamScorebook(path=None)
    r2_book.load_candidates([c])
    r2 = r2_book.rows[0]
    r2_book.update_row(r2, status="failed", error_code="begin_failed", error_message="x")
    r2_book.update_row(r2, status="recording", error_code=None, error_message=None)
    assert r2.status == "recording"


def test_scorebook_close_stops_writer_when_file_locked(tmp_path: Path) -> None:
    """close() 后即使 PermissionError 也不得让 writer 永久重试。"""
    import threading
    import time

    path = tmp_path / "locked.xlsx"
    calls = {"n": 0}

    def always_locked(_p, _rows):
        calls["n"] += 1
        raise PermissionError("locked by Excel")

    book = ExamScorebook(path=path, write_fn=always_locked, throttle_s=0.0)
    book.load_candidates([ExamCandidate(1, "A", "甲")])
    t0 = time.monotonic()
    book.close()
    elapsed = time.monotonic() - t0
    assert elapsed < 5.0
    worker = book._worker
    if worker is not None:
        worker.join(timeout=1.0)
        assert not worker.is_alive()
    # close 周期有限次尝试，不会无限涨
    n_after = calls["n"]
    time.sleep(0.6)
    assert calls["n"] == n_after


def test_student_id_zero_padded_number_format(tmp_path: Path) -> None:
    """数值 123 + 格式 000000 应按显示文本导入为 000123。"""
    from openpyxl import Workbook

    path = tmp_path / "pad.xlsx"
    wb = Workbook()
    ws = wb.active
    ws.append(["序号", "学号", "姓名"])
    ws.append([1, 123, "甲"])
    ws.cell(row=2, column=2).number_format = "000000"
    wb.save(path)
    wb.close()
    cands = import_roster_xlsx(path)
    assert cands[0].student_id == "000123"


def test_student_id_dashed_number_format(tmp_path: Path) -> None:
    """数值 123 + 格式 000-000 应按显示文本导入为 000-123。"""
    from openpyxl import Workbook

    path = tmp_path / "dash.xlsx"
    wb = Workbook()
    ws = wb.active
    ws.append(["序号", "学号", "姓名"])
    ws.append([1, 123, "甲"])
    ws.cell(row=2, column=2).number_format = "000-000"
    wb.save(path)
    wb.close()
    cands = import_roster_xlsx(path)
    assert cands[0].student_id == "000-123"


def test_student_id_escaped_and_quoted_formats(tmp_path: Path) -> None:
    """000\\-000 → 000-123；\"ID-\"000000 → ID-000123。"""
    from openpyxl import Workbook

    path = tmp_path / "esc.xlsx"
    wb = Workbook()
    ws = wb.active
    ws.append(["序号", "学号", "姓名"])
    ws.append([1, 123, "甲"])
    ws.cell(row=2, column=2).number_format = r"000\-000"
    wb.save(path)
    wb.close()
    cands = import_roster_xlsx(path)
    assert cands[0].student_id == "000-123"

    path2 = tmp_path / "quote.xlsx"
    wb = Workbook()
    ws = wb.active
    ws.append(["序号", "学号", "姓名"])
    ws.append([1, 123, "乙"])
    ws.cell(row=2, column=2).number_format = '"ID-"000000'
    wb.save(path2)
    wb.close()
    cands2 = import_roster_xlsx(path2)
    assert cands2[0].student_id == "ID-000123"


def test_student_id_unparseable_custom_format_rejects_sheet(tmp_path: Path) -> None:
    """无法还原的自定义格式不得静默裸数字，应整表拒绝。"""
    from openpyxl import Workbook

    path = tmp_path / "badfmt.xlsx"
    wb = Workbook()
    ws = wb.active
    ws.append(["序号", "学号", "姓名"])
    ws.append([1, 123, "甲"])
    # 科学计数等无法安全还原
    ws.cell(row=2, column=2).number_format = "0.00E+00"
    wb.save(path)
    wb.close()
    with pytest.raises(RosterError) as ei:
        import_roster_xlsx(path)
    assert ei.value.code == "student_id_not_text"


def test_failed_retest_does_not_mark_two_scoring_rows() -> None:
    """首次成功 + 补考失败：仅成功行计分。"""
    book = ExamScorebook(path=None)
    c = ExamCandidate(1, "S1", "张三")
    book.load_candidates([c])
    r1 = book.rows[0]
    book.update_row(
        r1,
        status="completed",
        front_score=0.8,
        side_score=0.8,
        combined_score=0.8,
        combined_percent=80,
    )
    r2 = book.append_retest(c)
    assert r2.is_scoring_row is False
    assert r1.is_scoring_row is True
    book.update_row(
        r2,
        status="failed",
        error_code="compare_failed",
        error_message="x",
    )
    assert r1.is_scoring_row is True
    assert r2.is_scoring_row is False
    assert r1.status == "completed"


def test_scorebook_flush_does_not_overwrite_with_stale_snapshot(tmp_path: Path) -> None:
    """旧快照写盘不得覆盖更新代（已完成成绩不丢）。"""
    import threading
    import time

    path = tmp_path / "scores.xlsx"
    write_gate = threading.Event()
    release_gate = threading.Event()
    writes: list[int] = []

    def slow_write(p: Path, rows: list) -> None:
        n = len(rows)
        writes.append(n)
        # 第一趟写故意卡住，模拟旧快照仍在落盘
        if len(writes) == 1:
            write_gate.set()
            release_gate.wait(timeout=5.0)
        write_scorebook_xlsx(p, rows)

    book = ExamScorebook(path=path, write_fn=slow_write, throttle_s=0.0)
    c1 = ExamCandidate(1, "A", "甲")
    c2 = ExamCandidate(2, "B", "乙")
    book.load_candidates([c1, c2])
    r1, r2 = book.rows

    # 后台 flush 拿 1 行完成快照
    book.update_row(
        r1,
        status="completed",
        front_score=0.8,
        side_score=0.8,
        combined_score=0.8,
        combined_percent=80,
    )
    t = threading.Thread(target=lambda: book.flush(), daemon=True)
    t.start()
    assert write_gate.wait(timeout=3.0)

    # 期间第二行完成并再 flush
    book.update_row(
        r2,
        status="completed",
        front_score=0.9,
        side_score=0.9,
        combined_score=0.9,
        combined_percent=90,
    )
    release_gate.set()
    t.join(timeout=5.0)
    book.flush()
    book.close()

    from openpyxl import load_workbook

    wb = load_workbook(path)
    ws = wb.active
    data_rows = list(ws.iter_rows(min_row=2, values_only=True))
    wb.close()
    assert len(data_rows) == 2
    # 两行综合分都应在文件中
    percents = {int(r[4]) for r in data_rows if r[4] not in (None, "")}
    assert percents == {80, 90}
